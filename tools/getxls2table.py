"""
getxls2table.py

将 assets/next1.xlsx 导入数据库表 company_addType。
用法：
    python tools/getxls2table.py
"""

from pathlib import Path

import openpyxl
import pymysql

# 项目根目录（tools 的上一级）
BASE_DIR = Path(__file__).resolve().parent.parent
ENV_MYSQL_PATH = BASE_DIR / "env" / "mysql"
XLSX_PATH = BASE_DIR / "assets" / "next1.xlsx"


# -------------------------------------------
#  数据库连接（参考 interface.py）
# -------------------------------------------


def load_mysql_config(path: Path = ENV_MYSQL_PATH) -> dict:
    """从 env/mysql 读取 KEY=VALUE 配置。"""
    config = {}
    with open(path, "r", encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if not line or line.startswith("#"):
                continue
            if "=" not in line:
                continue
            key, _, value = line.partition("=")
            config[key.strip()] = value.strip()
    return config


def get_connection():
    cfg = load_mysql_config()
    return pymysql.connect(
        host=cfg["MYSQL_HOST"],
        port=int(cfg.get("MYSQL_PORT", 3306)),
        user=cfg["MYSQL_USER"],
        password=cfg["MYSQL_PASSWORD"],
        database=cfg["MYSQL_DB"],
        charset="utf8mb4",
    )


# -------------------------------------------
#  建表 DDL
# -------------------------------------------

CREATE_TABLE_SQL = """
    CREATE TABLE IF NOT EXISTS company_addType (
        id BIGINT UNSIGNED NOT NULL AUTO_INCREMENT COMMENT '主键',
        external_company_id VARCHAR(64) DEFAULT '' COMMENT '外部公司ID',
        name VARCHAR(128) NOT NULL DEFAULT '' COMMENT '企业名称',
        region VARCHAR(64) DEFAULT '' COMMENT '所在区域',
        general_strength TEXT COMMENT '综合实力',
        main_business TEXT COMMENT '主营业务',
        corporate_vision TEXT COMMENT '企业愿景',
        product_list JSON COMMENT '产品列表',
        cooperation_case_list JSON COMMENT '合作案例列表',
        honor_cert_list JSON COMMENT '荣誉资质列表',
        status VARCHAR(20) NOT NULL DEFAULT 'success' COMMENT '状态',
        retry_count INT NOT NULL DEFAULT 0 COMMENT '重试次数',
        error_message TEXT COMMENT '错误信息',
        creator_id VARCHAR(36) NOT NULL DEFAULT '' COMMENT '创建者ID',
        create_time DATETIME COMMENT '创建时间',
        last_updater_id VARCHAR(36) NOT NULL DEFAULT '' COMMENT '最后更新者ID',
        last_update_time DATETIME COMMENT '最后更新时间',
        product_names TEXT COMMENT '产品列表（中文逗号分隔）',
        company_type VARCHAR(50) DEFAULT '' COMMENT '公司类型',
        classify_reason TEXT COMMENT '分类理由',
        PRIMARY KEY (id)
    ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COMMENT='企业分类表';
"""

# xlsx 列名 → 数据库列名（前17列同 aira_enterprise_info，后3列是新增）
COLUMN_MAPPING = [
    "id",
    "external_company_id",
    "name",
    "region",
    "general_strength",
    "main_business",
    "corporate_vision",
    "product_list",
    "cooperation_case_list",
    "honor_cert_list",
    "status",
    "retry_count",
    "error_message",
    "creator_id",
    "create_time",
    "last_updater_id",
    "last_update_time",
    "product_names",      # xlsx: 产品列表
    "company_type",       # xlsx: 公司类型
    "classify_reason",    # xlsx: 分类理由
]


# -------------------------------------------
#  主流程
# -------------------------------------------


def read_xlsx(path: Path):
    """读取 xlsx，返回 (列名列表, 行数据列表)。"""
    wb = openpyxl.load_workbook(path, read_only=True)
    ws = wb.active
    rows = []
    columns = None
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            columns = [str(c) if c else "" for c in row]
        else:
            rows.append(row)
    wb.close()
    return columns, rows


def main():
    print(f"读取 Excel: {XLSX_PATH}")
    xlsx_columns, rows = read_xlsx(XLSX_PATH)
    print(f"  sheet 列: {xlsx_columns}")
    print(f"  数据行数: {len(rows)}")

    # 检查列数对齐
    assert len(xlsx_columns) == len(COLUMN_MAPPING), (
        f"列数不匹配: xlsx {len(xlsx_columns)} != 映射 {len(COLUMN_MAPPING)}"
    )

    conn = get_connection()
    try:
        with conn.cursor() as cursor:
            # 建表
            cursor.execute(CREATE_TABLE_SQL)
            print("[OK] company_addType 表已就绪")

            # 清空旧数据
            cursor.execute("TRUNCATE TABLE company_addType")
            print("  已清空旧数据")

            # 批量插入
            placeholders = ", ".join(["%s"] * len(COLUMN_MAPPING))
            insert_sql = (
                f"INSERT INTO company_addType ({', '.join(COLUMN_MAPPING)}) "
                f"VALUES ({placeholders})"
            )

            # JSON 类型的列索引（product_list=7, cooperation_case_list=8, honor_cert_list=9，0-based）
            JSON_COL_INDEXES = {7, 8, 9}

            # 将 None 转为空串（JSON 列转为 '[]'），避免 NOT NULL 字段报错
            def safe_row(row):
                return tuple(
                    "[]" if v is None and i in JSON_COL_INDEXES
                    else "" if v is None
                    else v
                    for i, v in enumerate(row)
                )

            batch_size = 500
            total = len(rows)
            for start in range(0, total, batch_size):
                batch = [safe_row(r) for r in rows[start : start + batch_size]]
                cursor.executemany(insert_sql, batch)
                pct = min(start + batch_size, total) * 100 // total
                print(f"  进度: {min(start + batch_size, total)}/{total} ({pct}%)")

        conn.commit()
        print(f"[OK] 导入完成，共 {total} 条")

    finally:
        conn.close()


if __name__ == "__main__":
    main()
